#!/usr/bin/env python3
"""
Build the PowerSGD reproduction slide deck (PowerPoint, 16:9).

    python presentation/make_figures.py   # first, regenerates figures/
    python presentation/make_slides.py    # writes presentation/PowerSGD_reproduction.pptx
"""

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")
RESULTS = os.path.join(HERE, "..", "results", "paper_reproduction")
OUT = os.path.join(HERE, "PowerSGD_reproduction.pptx")

NAVY = RGBColor(0x1B, 0x2A, 0x41)
ACCENT = RGBColor(0x2E, 0x86, 0xAB)
GRAY = RGBColor(0x5C, 0x67, 0x7D)
LIGHT = RGBColor(0xEE, 0xF3, 0xF7)
GREEN = RGBColor(0x3A, 0x9D, 0x5D)
RED = RGBColor(0xC8, 0x55, 0x3D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W, SLIDE_H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
BLANK = prs.slide_layouts[6]


def text_box(slide, left, top, width, height, text, size=18, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
    return box


def bullets(slide, left, top, width, height, items, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for i, item in enumerate(items):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(8)
        bullet = paragraph.add_run()
        bullet.text = "•  "
        bullet.font.size = Pt(size)
        bullet.font.color.rgb = ACCENT
        bullet.font.bold = True
        # "**bold part** rest" renders the leading part in bold
        if item.startswith("**"):
            head, rest = item[2:].split("**", 1)
            run = paragraph.add_run()
            run.text = head
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = NAVY
            item = rest
        run = paragraph.add_run()
        run.text = item
        run.font.size = Pt(size)
        run.font.color.rgb = NAVY
    return box


def picture(slide, path, top, max_width, max_height, left=None):
    """Add an image scaled to fit max_width x max_height, centered horizontally by default."""
    with Image.open(path) as image:
        aspect = image.width / image.height
    width = min(max_width, max_height * aspect)
    height = width / aspect
    if left is None:
        left = (SLIDE_W - width) / 2
    slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(width), Inches(height))
    return top + height


def table(slide, left, top, width, rows, col_widths, size=14, row_height=0.42, first_col_bold=True):
    shape = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width),
        Inches(row_height * len(rows)),
    )
    tbl = shape.table
    for j, w in enumerate(col_widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = ""
            paragraph = cell.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = str(value)
            run.font.size = Pt(size)
            if i == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
                run.font.bold = True
                run.font.color.rgb = WHITE
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
                run.font.color.rgb = NAVY
                run.font.bold = first_col_bold and j == 0
                if value.startswith("✓"):
                    run.font.color.rgb = GREEN
                    run.font.bold = True
                elif value.startswith("✗"):
                    run.font.color.rgb = RED
                    run.font.bold = True
    return shape


def new_slide(title, takeaway=None, notes=None):
    slide = prs.slides.add_slide(BLANK)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), Inches(SLIDE_H))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    text_box(slide, 0.6, 0.3, 12.2, 0.8, title, size=30, bold=True)
    if takeaway:
        text_box(slide, 0.6, 1.0, 12.2, 0.6, takeaway, size=19, color=ACCENT, bold=True)
    number = len(prs.slides)
    text_box(slide, 11.9, 7.0, 1.1, 0.4, str(number), size=12, color=GRAY, align=PP_ALIGN.RIGHT)
    text_box(slide, 0.6, 7.0, 8, 0.4, "Reproducing PowerSGD on Perlmutter", size=12, color=GRAY)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


# 1. Title
slide = prs.slides.add_slide(BLANK)
band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
band.fill.solid()
band.fill.fore_color.rgb = NAVY
band.line.fill.background()
stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.55), Inches(2.2), Inches(0.08))
stripe.fill.solid()
stripe.fill.fore_color.rgb = ACCENT
stripe.line.fill.background()
text_box(slide, 0.9, 1.7, 11.5, 1.2, "Reproducing PowerSGD on Perlmutter", size=44, color=WHITE, bold=True)
text_box(
    slide, 0.9, 2.75, 11.5, 0.8,
    "Low-rank gradient compression for data-parallel training",
    size=24, color=RGBColor(0xC9, 0xD6, 0xE3),
)
text_box(
    slide, 0.9, 3.9, 11.5, 0.6,
    "Paper: Vogels, Karimireddy & Jaggi, “PowerSGD: Practical Low-Rank Gradient Compression "
    "for Distributed Optimization”, NeurIPS 2019",
    size=16, color=RGBColor(0xC9, 0xD6, 0xE3), italic=True,
)
text_box(slide, 0.9, 5.6, 11.5, 0.5, "Nayan Yadav", size=22, color=WHITE, bold=True)
text_box(slide, 0.9, 6.1, 11.5, 0.5, "Deep Learning course project  ·  November 2026",
         size=16, color=RGBColor(0xC9, 0xD6, 0xE3))
slide.notes_slide.notes_text_frame.text = (
    "We reproduced the main results of the PowerSGD paper on NERSC Perlmutter, using the "
    "authors' own research code, and checked which of its claims still hold on modern hardware."
)

# 2. Problem
slide = new_slide(
    "The problem: communication in data-parallel training",
    "Every training step, every GPU must exchange its entire gradient.",
    "Data parallelism splits the batch across GPUs, but the gradients must be averaged after every "
    "step. The message is as large as the model, so communication grows with model size.",
)
bullets(slide, 0.6, 1.9, 12.0, 4.8, [
    "**Data parallelism: **each GPU computes a gradient on its share of the batch; all GPUs then "
    "average their gradients (all-reduce) before updating the model.",
    "**The message is the whole model: **43 MB per step for ResNet18 and 110 MB for the 3-layer "
    "LSTM used in the paper.",
    "**As models and GPU counts grow, **communication can dominate the time per step.",
    "**Earlier compressors **(sign-based, top-K) either lose accuracy or cannot be summed with "
    "all-reduce, so they scale poorly with more GPUs.",
], size=20)

# 3. Method
slide = new_slide(
    "PowerSGD: send a low-rank approximation of each gradient",
    "Approximate each layer's gradient matrix M by P·Qᵀ with small rank r (1, 2 or 4).",
    "Each layer's gradient is a matrix. PowerSGD runs one step of power iteration to find a rank-r "
    "approximation and only communicates the two thin factors. Warm start and error feedback keep "
    "it accurate; linearity lets it use all-reduce.",
)
bottom = picture(slide, os.path.join(FIG, "fig1_method.png"), 1.7, 11.0, 2.9)
bullets(slide, 0.6, bottom + 0.2, 12.0, 2.4, [
    "**One power-iteration step per layer: **P = M·Q → all-reduce P → orthogonalize P → "
    "Q = Mᵀ·P → all-reduce Q.",
    "**Warm start: **reuse the previous step's Q, so a single cheap iteration is enough.",
    "**Error feedback: **remember what the compression missed and add it to the next gradient.",
    "**Linear in M: **the GPUs' P and Q can simply be summed with all-reduce.",
], size=18)

# 4. Claims
slide = new_slide(
    "The paper's three claims",
    "We tested all three with the authors' own code, matching Table 3 of the paper.",
    "These are the three claims we set out to check. Numbers are from Table 3 of the paper, "
    "16 GPUs, ResNet18 on CIFAR-10 and an LSTM on WikiText-2.",
)
claims = [
    ("1", "Same accuracy as SGD",
     "CIFAR-10: Rank 2 reaches 94.4% vs 94.3% for SGD.\n\nWikiText-2: Rank 4 matches SGD's "
     "perplexity (91 vs 91)."),
    ("2", "100× or more less communication",
     "136× less data with Rank 2 on ResNet18.\n\n120× less with Rank 4 on the LSTM."),
    ("3", "Faster training",
     "On 16 GPUs, time per batch drops by 23% (ResNet18, Rank 2) and 55% (LSTM, Rank 4)."),
]
for k, (number, head, body) in enumerate(claims):
    left = 0.6 + k * 4.15
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.9),
                                  Inches(3.85), Inches(4.6))
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.fill.background()
    text_box(slide, left + 0.3, 2.1, 3.3, 0.8, number, size=40, color=ACCENT, bold=True)
    text_box(slide, left + 0.3, 2.95, 3.3, 0.9, head, size=22, bold=True)
    text_box(slide, left + 0.3, 3.95, 3.3, 2.4, body.split("\n"), size=17, color=GRAY)

# 5. Setup
slide = new_slide(
    "Setup: same code and hyperparameters, very different hardware",
    "Only the hardware changed; the algorithm, models and hyperparameters are the paper's.",
    "We ran the authors' research code unchanged except for porting. The big difference is the "
    "network: the paper's 10 Gbit/s Ethernet versus Perlmutter's NVLink and Slingshot.",
)
table(slide, 0.6, 1.8, 12.1, [
    ["", "Paper (2019)", "Ours: Nayan Yadav's runs on NERSC Perlmutter (2026)"],
    ["GPUs", "16× GTX Titan X (2 per node, 8 nodes)", "16× A100 (4 per node, 4 nodes)"],
    ["Network", "10 Gbit/s Ethernet", "NVLink within a node, Slingshot-11 between nodes"],
    ["Software", "PyTorch 1.1, NCCL", "PyTorch 2.13, NCCL"],
    ["Tasks", "ResNet18 / CIFAR-10 (300 epochs), LSTM / WikiText-2 (90 epochs)", "Same"],
    ["Hyperparameters", "Tuned for SGD, reused unchanged for PowerSGD", "Identical"],
    ["Repetitions", "3 seeds", "3 seeds per method: 24 runs"],
    ["Compute used", "Not reported", "36.9 GPU-hours (from Slurm accounting)"],
], col_widths=[2.2, 4.9, 5.0], size=15, row_height=0.5)
bullets(slide, 0.6, 5.85, 12.0, 1.1, [
    "**Porting only: **one GPU per process under Slurm; replaced two libraries that no longer exist "
    "(a CUDA bit-packing extension and the old torchtext/spaCy text pipeline); fixed the dataset "
    "download.",
], size=16)

# 6. Claim 1
slide = new_slide(
    "Claim 1: accuracy matches SGD  ✓",
    "CIFAR-10 is within 0.3 points of the paper; LSTM perplexity is 8-9 points higher for "
    "every method.",
    "Left: CIFAR-10 test accuracy. Rank 2 and Rank 4 hit the paper's numbers exactly. Right: the "
    "LSTM. All of our perplexities are higher than the paper's, including plain SGD, which points "
    "to a data difference rather than PowerSGD. The next slides show this.",
)
bottom = picture(slide, os.path.join(FIG, "fig2_quality.png"), 1.7, 12.0, 4.35)
bullets(slide, 0.6, bottom + 0.1, 12.0, 1.0, [
    "Seeds agree closely (error bars = min-max of 3 runs), so the small differences are real "
    "but tiny.",
], size=16)

# 7. CIFAR curves
slide = new_slide(
    "CIFAR-10 training curves",
    "Rank 2 and Rank 4 track SGD; Rank 1 levels off about 1 point lower, as in the paper.",
    "These are the counterpart of the paper's Figure 4. The jump at epoch 150 is the learning-rate "
    "decay. Training took about 8 minutes on 16 A100s versus about 45 minutes in the paper.",
)
bottom = picture(slide, os.path.join(RESULTS, "curves_cifar16_cifar.png"), 1.7, 12.2, 4.6)
text_box(slide, 0.6, bottom + 0.1, 12.0, 0.6,
         "Jump at epoch 150 = learning-rate decay. Shaded = min-max over 3 seeds. "
         "Full run: ~8 minutes on 16 A100s (paper: ~45 minutes).",
         size=15, color=GRAY)

# 8. LSTM curves
slide = new_slide(
    "WikiText-2 LSTM training curves",
    "Higher rank gives lower perplexity; Rank 4 ends level with SGD, as in the paper.",
    "Same ordering as the paper's Figure 4: Rank 1 is clearly worse, Rank 2 is close and Rank 4 "
    "matches SGD. The drop at epoch 60 is the learning-rate decay.",
)
bottom = picture(slide, os.path.join(RESULTS, "curves_lm16_lm.png"), 1.7, 12.2, 4.6)
text_box(slide, 0.6, bottom + 0.1, 12.0, 0.6,
         "Drop at epoch 60 = learning-rate decay. Shaded = min-max over 3 seeds. "
         "Full run: ~3 minutes on 16 A100s.",
         size=15, color=GRAY)

# 9. Gap to SGD
slide = new_slide(
    "The key comparison: PowerSGD relative to SGD",
    "Measured against SGD, our results match the paper almost exactly, LSTM included.",
    "The paper's claim is about PowerSGD compared with SGD, so this is the fair comparison. Our gaps "
    "match the paper's within half a point of accuracy and about one point of perplexity. The "
    "LSTM's constant offset cancels out because it also affects SGD.",
)
bottom = picture(slide, os.path.join(FIG, "fig3_gap_to_sgd.png"), 1.7, 12.0, 4.3)
bullets(slide, 0.6, bottom + 0.1, 12.0, 1.1, [
    "**The LSTM's +8-9 perplexity offset also affects SGD, **so it comes from the data pipeline "
    "(the original spaCy tokenizer had to be replaced), not from PowerSGD.",
], size=16)

# 10. Claim 2
slide = new_slide(
    "Claim 2: 100× or more less communication  ✓",
    "Data sent per epoch matches the paper for every method.",
    "Log scale. Communication volume is a property of the algorithm and the model shapes, so it "
    "should match exactly, and it does. Rank 2 sends 136 times less than SGD on ResNet18.",
)
bottom = picture(slide, os.path.join(FIG, "fig4_communication.png"), 1.7, 12.0, 4.3)
bullets(slide, 0.6, bottom + 0.1, 12.0, 1.1, [
    "**ResNet18: **Rank 2 sends 136× less than SGD.   **LSTM: **Rank 1 sends 310× less, "
    "Rank 4 sends 120× less.",
], size=16)

# 11. Claim 3
slide = new_slide(
    "Claim 3: faster training  ✗ on Perlmutter",
    "On ResNet18, PowerSGD is slower than SGD; on the LSTM it saves only 3-8%.",
    "This is where our results differ from the paper. We compare the relative change against SGD on "
    "the same hardware, not raw milliseconds, since A100s are simply faster than Titan X GPUs.",
)
bottom = picture(slide, os.path.join(FIG, "fig5_speed.png"), 1.7, 12.0, 4.3)
bullets(slide, 0.6, bottom + 0.05, 12.2, 1.2, [
    "**Paper: **16-26% faster (ResNet18) and 53-56% faster (LSTM).   **Ours: **8-27% slower "
    "and 3-8% faster.",
    "**Our ms per batch: **ResNet18 SGD 26.2 → Rank 1/2/4: 28.3 / 29.6 / 33.2.   "
    "LSTM SGD 28.8 → 26.5 / 27.1 / 28.1.",
], size=15)

# 12. Why
slide = new_slide(
    "Why: communication is cheap on Perlmutter",
    "There is little communication left to save, and compressing costs more than it saves.",
    "Breakdown of one training step. On CIFAR-10, SGD's whole gradient exchange is only 6.4 ms, "
    "about a quarter of the step, which caps any possible gain. PowerSGD's compression and "
    "communication take longer than that, because they are many small operations limited by "
    "latency rather than bandwidth.",
)
bottom = picture(slide, os.path.join(FIG, "fig6_breakdown.png"), 1.6, 9.0, 3.4)
bullets(slide, 0.6, bottom + 0.1, 12.2, 2.0, [
    "**SGD's full all-reduce is only 6.4 of 26 ms (CIFAR-10) and 4.4 of 29 ms (LSTM): **"
    "the most any compressor could save.",
    "**PowerSGD's step takes longer (9.2 ms at Rank 2): **many small per-layer matrix products and "
    "orthogonalizations, plus small all-reduces bound by latency, not bandwidth.",
    "**The paper's 10 Gbit/s network is ~80× slower **than Perlmutter's 4×200 Gbit/s links between "
    "nodes, so there communication dominated.",
], size=15)

# 13. Scorecard
slide = new_slide(
    "Scorecard",
    "Accuracy and compression are properties of the algorithm; speedup depends on the hardware.",
    "Summary of what reproduced. The algorithm behaves exactly as described; the wall-clock benefit "
    "depends on how expensive communication is relative to computation.",
)
table(slide, 0.6, 1.85, 12.1, [
    ["Claim", "Paper", "Ours", "Reproduced?"],
    ["Accuracy on par with SGD (CIFAR-10)", "Rank 2: 94.4% vs 94.3%", "Rank 2: 94.4% vs 94.1%", "✓ Yes"],
    ["LSTM: Rank 4 matches SGD", "91 vs 91 perplexity", "98.9 vs 99.6 perplexity",
     "✓ Yes (relative)"],
    ["Quality improves with rank", "Rank 1 < 2 < 4", "Rank 1 < 2 < 4 (both tasks)", "✓ Yes"],
    ["Communication reduction", "136× (Rank 2, ResNet18)", "Same MB per epoch", "✓ Yes"],
    ["Speedup, ResNet18", "−23% per batch (Rank 2)", "+13% per batch", "✗ No"],
    ["Speedup, LSTM", "−55% per batch (Rank 4)", "−3% per batch", "✗ Mostly not"],
], col_widths=[3.9, 3.0, 3.1, 2.1], size=16, row_height=0.58)

# 14. Takeaways and next steps
slide = new_slide(
    "Takeaways and next steps",
    notes="PowerSGD's speedup depends on the communication-to-compute ratio. Larger models such as "
          "GPT-style transformers send much bigger gradients, which is where we expect compression to "
          "pay off on modern hardware. That is the next part of the project.",
)
text_box(slide, 0.6, 1.2, 5.8, 0.5, "What we found", size=22, color=ACCENT, bold=True)
bullets(slide, 0.6, 1.8, 5.9, 4.8, [
    "The paper's accuracy and compression results reproduce from the authors' code, with no "
    "retuning: 24/24 runs, 3 seeds each.",
    "Its speedups do not: on a fast interconnect, communication is too small a part of each step "
    "for compression to pay off on these models.",
    "Fair comparisons across hardware use changes relative to SGD, not raw times.",
], size=18)
text_box(slide, 6.9, 1.2, 5.8, 0.5, "Next steps", size=22, color=ACCENT, bold=True)
bullets(slide, 6.9, 1.8, 5.9, 4.8, [
    "**nanoGPT and larger models: **much bigger gradients, so communication matters more.",
    "**Data vs tensor parallelism: **where gradient compression applies, and where it does not.",
    "**Low-precision gradients (fp32 → fp8): **an alternative or complement to low-rank "
    "compression.",
    "**Slow-network runs **(Gloo backend, 1-16 GPUs) to recreate the paper's setting.",
], size=18)

# 15. Appendix
slide = new_slide(
    "Appendix: full Table 3 results",
    "Nayan Yadav's Perlmutter runs: 16 GPUs, 3 seeds each, shown as mean [min, max].",
    "All numbers behind the charts. Generated with perlmutter/submit_paper.sh table3 and "
    "paper-code/analyze_results.py.",
)
table(slide, 0.6, 1.75, 12.1, [
    ["Task", "Method", "Ours", "Paper", "MB/epoch ours", "Paper", "ms/batch ours", "Paper"],
    ["CIFAR-10 (acc.)", "SGD", "94.1% [94.1, 94.1]", "94.3%", "1,073", "1,023", "26.2", "312"],
    ["", "Rank 1", "93.3% [93.1, 93.4]", "93.6%", "4", "4", "28.3", "229"],
    ["", "Rank 2", "94.4% [94.3, 94.4]", "94.4%", "8", "8", "29.6", "239"],
    ["", "Rank 4", "94.5% [94.3, 94.7]", "94.5%", "15", "14", "33.2", "260"],
    ["WikiText-2 (ppl.)", "SGD", "99.6 [99.4, 99.9]", "91", "7,882", "7,730", "28.8", "300"],
    ["", "Rank 1", "111.2 [110.5, 112.0]", "102", "25", "25", "26.5", "131"],
    ["", "Rank 2", "101.4 [100.9, 102.1]", "93", "39", "38", "27.1", "141"],
    ["", "Rank 4", "98.9 [98.3, 99.3]", "91", "65", "64", "28.1", "134"],
], col_widths=[2.0, 1.1, 2.4, 1.1, 1.6, 1.2, 1.6, 1.1], size=14, row_height=0.5)
text_box(slide, 0.6, 6.45, 12.1, 0.5,
         "CIFAR-10 SGD seeds: 94.075, 94.075, 94.083 (different training runs that happen to end "
         "equal). Raw ms/batch is not comparable across hardware; compare the change vs SGD.",
         size=14, color=GRAY, italic=True)

prs.save(OUT)
print("saved", OUT, f"({len(prs.slides)} slides)")
